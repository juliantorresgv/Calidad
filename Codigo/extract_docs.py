from pathlib import Path
from docx import Document
from openpyxl import load_workbook

def extract_docx(p: Path) -> str:
    doc = Document(str(p))
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    tables = []
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells]
            if any(row_text):
                tables.append(" | ".join(row_text))
    return "\n".join(paragraphs + tables)

def extract_xlsx(p: Path) -> str:
    wb = load_workbook(str(p), data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"--- Hoja: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() for cell in row):
                lines.append(" | ".join(str(cell) for cell in row))
    return "\n".join(lines)

def main():
    docs_dir = Path(r"C:\Users\1121871773\OneDrive - agvco\Documentos\Agente Calidad Codigo\Documentación")
    out = Path(r"C:\Users\1121871773\OneDrive - agvco\Documentos\Agente Calidad Codigo\Codigo\docs_text.txt")
    with out.open("w", encoding="utf-8", errors="ignore") as f:
        for p in sorted(docs_dir.iterdir()):
            if p.suffix.lower() == ".docx":
                f.write(f"\n{'='*60}\nARCHIVO: {p.name}\n{'='*60}\n")
                try:
                    f.write(extract_docx(p))
                except Exception as e:
                    f.write(f"ERROR leyendo {p.name}: {e}")
            elif p.suffix.lower() == ".xlsx":
                f.write(f"\n{'='*60}\nARCHIVO: {p.name}\n{'='*60}\n")
                try:
                    f.write(extract_xlsx(p))
                except Exception as e:
                    f.write(f"ERROR leyendo {p.name}: {e}")
    print(f"Texto extraído en {out}")

if __name__ == "__main__":
    main()
